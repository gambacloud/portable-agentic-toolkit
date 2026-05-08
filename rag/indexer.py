"""Parse files, chunk text, embed, and store in ChromaDB."""
from __future__ import annotations

import csv
import io
import re
from pathlib import Path

from utils.logger import get_logger

log = get_logger(__name__)

SUPPORTED_SUFFIXES = {".txt", ".md", ".pdf", ".docx", ".csv", ".xlsx", ".xls"}


def parse_file(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix in (".txt", ".md"):
        return path.read_text(encoding="utf-8", errors="replace")
    if suffix == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(str(path))
        pages = [page.extract_text() or "" for page in reader.pages]
        return "\n\n".join(p for p in pages if p.strip())
    if suffix == ".docx":
        from docx import Document
        doc = Document(str(path))
        return "\n\n".join(p.text for p in doc.paragraphs if p.text.strip())
    if suffix == ".csv":
        text = path.read_text(encoding="utf-8", errors="replace")
        reader = csv.reader(io.StringIO(text))
        rows = ["\t".join(row) for row in reader]
        return "\n".join(rows)
    if suffix in (".xlsx", ".xls"):
        import openpyxl
        wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
        sheets = []
        for ws in wb.worksheets:
            rows = ["\t".join(str(c.value or "") for c in row) for row in ws.iter_rows()]
            sheets.append(f"[{ws.title}]\n" + "\n".join(rows))
        return "\n\n".join(sheets)
    # fallback
    return path.read_text(encoding="utf-8", errors="replace")


def chunk_text(text: str, size: int = 1500, overlap: int = 150) -> list[str]:
    paras = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
    chunks: list[str] = []
    current = ""
    for para in paras:
        if len(current) + len(para) + 2 <= size:
            current = (current + "\n\n" + para).strip() if current else para
        else:
            if current:
                chunks.append(current)
            if len(para) > size:
                for i in range(0, len(para), size - overlap):
                    piece = para[i : i + size]
                    if piece.strip():
                        chunks.append(piece)
                current = ""
            else:
                current = para
    if current:
        chunks.append(current)
    return chunks


def index_file(path: Path, collection) -> int:
    from rag.embedder import embed

    log.info("Indexing %s", path.name)
    text = parse_file(path)
    if not text.strip():
        log.warning("No text extracted from %s", path.name)
        return 0

    chunks = chunk_text(text)
    if not chunks:
        return 0

    source = path.name
    # Remove any existing chunks for this source
    try:
        existing = collection.get(where={"source": source})
        if existing["ids"]:
            collection.delete(ids=existing["ids"])
            log.debug("Removed %d old chunks for %s", len(existing["ids"]), source)
    except Exception:
        pass

    embeddings = embed(chunks)
    ids = [f"{source}::chunk{i}" for i in range(len(chunks))]
    collection.add(
        documents=chunks,
        embeddings=embeddings,
        ids=ids,
        metadatas=[{"source": source, "chunk": i} for i in range(len(chunks))],
    )
    log.info("Indexed %s — %d chunks", source, len(chunks))
    return len(chunks)
